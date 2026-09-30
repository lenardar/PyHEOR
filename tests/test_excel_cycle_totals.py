"""Evaluate the generated scalar formula subset independently of the engines.

This checks references and arithmetic, not a native Excel recalculation.
"""
import ast
import math
import operator
import re
import pytest
from openpyxl import load_workbook
import pyheor as ph


def calculate(ws, address):
    def walk(node):
        if isinstance(node, ast.Constant): return node.value
        if isinstance(node, ast.BinOp):
            operation={ast.Add:operator.add, ast.Sub:operator.sub, ast.Mult:operator.mul,
                       ast.Div:operator.truediv, ast.Pow:operator.pow}[type(node.op)]
            return operation(walk(node.left),walk(node.right))
        if isinstance(node, ast.UnaryOp):
            return -walk(node.operand) if isinstance(node.op,ast.USub) else walk(node.operand)
        if isinstance(node, ast.Compare):
            assert len(node.ops)==1 and isinstance(node.ops[0],ast.Eq)
            return walk(node.left)==walk(node.comparators[0])
        if isinstance(node, ast.Call):
            name=node.func.id
            if name=='CELL': return calculate(ws, node.args[0].value)
            if name=='IF': return walk(node.args[1] if walk(node.args[0]) else node.args[2])
            if name=='EXP': return math.exp(walk(node.args[0]))
        raise AssertionError(ast.dump(node))
    value=ws[address].value
    if not isinstance(value,str) or not value.startswith('='): return value or 0
    expression=value[1:].replace('$','').replace('^','**')
    expression=re.sub(r'\b[A-Z]+[1-9][0-9]*\b',lambda m: f'CELL("{m.group()}")',expression)
    expression=re.sub(r'(?<![<>=!])=(?!=)','==',expression)
    return walk(ast.parse(expression,mode='eval').body)


@pytest.mark.parametrize('method',['beginning','end','life-table'])
@pytest.mark.parametrize('kind',['markov','psm','terminal'])
def test_export_totals_follow_cycle_reward_semantics(tmp_path, method, kind):
    cycle=ph.Cycle(3,'week')
    kwargs=dict(strategies=['S'],n_cycles=3,cycle=cycle,method=method,
                dr_cost=lambda p: ph.rescale_discount_rate(.045,ph.Cycle(1,'year'),cycle),dr_qaly=.01)
    if kind=='markov':
        model=ph.MarkovModel(states=['Alive','Dead'],**kwargs)
        model.set_transitions('S', lambda p,k: [[1-.1*k,.1*k],[0,1]])
        model.set_transition_cost('event','Alive','Dead',lambda p,k: 500/k)
        model.set_entry_qaly('entry','Alive',-.02)
        model.set_entry_cost('death','Dead',40)
    else:
        states=['PFS','PD','Terminal','Dead'] if kind=='terminal' else ['PFS','PD','Dead']
        model=ph.PSMModel(states=states,survival_endpoints=['PFS','OS'],
                         terminal_state='Terminal' if kind=='terminal' else None,**kwargs)
        model.set_survival('S','PFS',ph.rescale_survival(ph.Exponential(.4),from_unit='month',to_period=cycle))
        model.set_survival('S','OS',ph.rescale_survival(ph.Exponential(.1),from_unit='month',to_period=cycle))
        if kind=='terminal': model.set_state_cost('terminal',{'Terminal':200})
    alive=model.states[0]
    model.set_state_cost('care',{alive:lambda p,k: 100*k})
    model.set_state_qaly('health',{alive:ph.qaly(.8,cycle)})
    model.set_state_cost('limited',{alive:20},cycles=[1,3])
    model.set_starting_cost('start',75)
    model.set_starting_qaly('loss',-.01)
    expected=model.run_base_case().results['S']
    path=tmp_path/'audit.xlsx'
    ph.export_excel_model(model,path)
    ws=load_workbook(path)['Calc_S']
    assert calculate(ws,'B7')==pytest.approx(sum(expected['total_costs'].values()),rel=1e-12)
    assert calculate(ws,'B8')==pytest.approx(expected['total_qalys'],rel=1e-12)
    # Editable rewards propagate through formula references.
    row=next(row[0].row for row in ws.iter_rows() if row[0].value=='Starting cost / start')
    ws.cell(row,2,175)
    assert calculate(ws,'B7')==pytest.approx(sum(expected['total_costs'].values())+100)
