"""Cross-engine reward units, reconciliation and result workbook detail."""
import numpy as np
import pytest
from openpyxl import load_workbook
import pyheor as ph


def build(kind):
    cycle=ph.Cycle(1,'month')
    shared=dict(states=['Alive','Dead'],strategies=['SOC','TRT'],dr_cost=.01,dr_qaly=.02)
    if kind=='des':
        m=ph.DESModel(**shared,time_horizon=2,time_unit='month')
        period=m.time_period
    else:
        shared.update(n_cycles=2,cycle=cycle)
        if kind=='psm':
            m=ph.PSMModel(**shared,survival_endpoints=['OS'])
            for arm in m.strategy_names:m.set_survival(arm,'OS',ph.KaplanMeier([0],[1]))
        else:
            m=ph.MicroSimModel(**shared,n_patients=3) if kind=='micro' else ph.MarkovModel(**shared)
            for arm in m.strategy_names:m.set_transitions(arm,[[1,0],[0,1]])
        period=cycle
    m.set_state_cost('care',{'Alive':100})
    m.set_state_qaly('health',{'Alive':ph.qaly(.8,period)})
    m.set_starting_cost('initial',{'SOC':20,'TRT':30})
    m.set_starting_qaly('ae',-.01)
    if kind!='psm':m.set_entry_cost('entry','Alive',50)
    if kind=='des':return m.run(n_patients=3,seed=42,progress=False)
    if kind=='micro':return m.run_base_case(seed=42,verbose=False)
    return m.run_base_case()


@pytest.mark.parametrize('kind',['markov','psm','micro','des'])
def test_components_reconcile_with_totals_and_workbooks(tmp_path,kind):
    result=build(kind)
    table=result.reward_components
    for arm in result.model.strategy_names:
        rows=table[table.Strategy==arm]
        assert rows[rows.Reward=='cost'].Discounted.sum()==pytest.approx(result._totals(arm)[0])
        assert rows[rows.Reward=='qaly'].Discounted.sum()==pytest.approx(result._totals(arm)[1])
        assert rows[rows.Category=='health'].Undiscounted.iloc[0]==pytest.approx(.8*2/12)
        assert rows[rows.Category=='ae'].Undiscounted.iloc[0]==-.01
    assert result.metadata['QALYUnit']=='QALY'
    path=tmp_path/(kind+'.xlsx')
    ph.export_to_excel(result,path)
    wb=load_workbook(path)
    assert {'Reward Components','Calculation Metadata'}<=set(wb.sheetnames)
    assert wb['Reward Components'].max_row==len(table)+1
    if kind=='des':
        assert result.time_in_state.TimeUnit.unique().tolist()==['month']
        assert result.survival_curve().TimeUnit.unique().tolist()==['month']
        assert {'Event Log','Time in State','Patient Outcomes'}<=set(wb.sheetnames)
        with pytest.raises(ValueError,match='continuous'): _=result.cycle_rewards
    else:
        cycles=result.cycle_rewards
        assert set(cycles.Cycle)=={1,2}
        for arm in result.model.strategy_names:
            for reward in ('cost','qaly'):
                detail=cycles[(cycles.Strategy==arm)&(cycles.Reward==reward)]
                total=table[(table.Strategy==arm)&(table.Reward==reward)]
                assert detail.Discounted.sum()==pytest.approx(total.Discounted.sum())
                np.testing.assert_allclose(detail.Discounted,detail.Undiscounted*detail.DiscountFactor)
        assert {'Cycle Rewards','State Occupancy'}<=set(wb.sheetnames)


def test_result_tables_preserve_original_calculation_settings():
    result=build('markov')
    before=result.cycle_rewards.copy()
    result.model.dr_cost=.5
    result.model.method='end'
    result.model.cycle=ph.Cycle(1,'year')
    assert result.metadata['CycleUnit']=='month'
    assert result.metadata['Method']=='life-table'
    assert result.metadata['CostDiscountRate']==.01
    assert result.cycle_rewards.equals(before)
    assert result.state_occupancy.Method.unique().tolist()==['life-table']


def test_des_censoring_passes_attributes_to_qaly_callbacks():
    m=ph.DESModel(states=['Alive','Dead'],strategies=['SOC'],time_horizon=2,time_unit='month')
    m.set_event('SOC','Alive','Dead',ph.Exponential(rate=1e-100))
    m.set_state_qaly('health',{'Alive':lambda p,t,a:ph.qaly(a['u'],m.time_period)})
    result=m.run(n_patients=2,seed=42,progress=False,attrs={'u':np.array([.6,.9])})
    np.testing.assert_allclose(result.patient_outcomes['Total QALYs'],[.6*2/12,.9*2/12])
    assert result.reward_components.Discounted.iloc[0]==pytest.approx(.75*2/12)


@pytest.mark.parametrize('kind',['micro','des'])
def test_individual_psa_export(tmp_path,kind):
    m=build(kind).model
    psa=m.run_psa(n_outer=2,n_inner=2,seed=1,verbose=False) if kind=='micro' else m.run_psa(n_sim=2,n_patients=2,seed=1,progress=False)
    path=tmp_path/'psa.xlsx'
    ph.export_to_excel(psa,path,include_psa=True)
    assert {'PSA Summary','CE Table','CEAC Data','Sampled Parameters'}<=set(load_workbook(path).sheetnames)
